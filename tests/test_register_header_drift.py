"""A register whose header has drifted is never imported column-shifted (AUD-041).

ADR 0012 names the failure: «an imported year whose columns are one to the
left». The header check always *found* it — the planner printed the finding and
then planned every row of the sheet as if the columns matched. Now a contracted
column whose header the sheet does not carry fails closed everywhere a sheet is
read to write or compare:

* `import_legacy_register` — every matter row of that sheet waits for review
  with the blocking `PARSE_ERROR`; apply refuses, and with
  `--accept-review-rows` those rows are left out, never guessed;
* the 2026 Excel pilot — the reader refuses before a row is read;
* the snapshot delta — the comparison stops as ambiguous.

A column the contract does not describe (`uncontracted`) is still only a
finding: the 2026 sheet carries two such columns legitimately.

The fixture inserts one column in the middle of the 2026 sheet, the way a
person adds «UUS VEERG» beside KELLELE: everything to its right moves.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from app.legacy_import.enums import Anomaly, RowOutcome
from app.legacy_import.excel_pilot import PilotError, read_pilot_sheet
from app.legacy_import.planner import HEADER_DRIFT_NOTE, build_plan
from app.legacy_import.snapshot_delta import DeltaRefused
from app.legacy_import.snapshot_delta import build_report as build_delta
from app.matters.models import Matter
from tests.synthetic_register import Row, Sheet, write_workbook

pytestmark = pytest.mark.django_db


def _register(tmp_path: Path) -> Path:
    return write_workbook(
        tmp_path / "register.xlsx",
        [
            Sheet(
                2026,
                [
                    Row(reference="2026_1", title="Sünteetiline eelnõu üks", owner="Mari"),
                    Row(reference="2026_2", title="Sünteetiline eelnõu kaks", owner="Mari"),
                ],
            ),
            Sheet(2025, [Row(reference="2025_1", title="Eelmise aasta teema", owner="Mari")]),
        ],
    )


def _insert_a_column(path: Path, *, at: int = 4, header: str = "UUS VEERG") -> Path:
    """A person inserts a column: every header to its right shifts by one."""
    workbook = load_workbook(path)
    sheet = workbook["2026"]
    sheet.insert_cols(at)
    sheet.cell(row=1, column=at).value = header
    workbook.save(path)
    workbook.close()
    return path


def _append_a_column(path: Path, header: str = "MÄRKUSED") -> Path:
    """A labelled column after the contract's last one: shifts nothing."""
    workbook = load_workbook(path)
    sheet = workbook["2026"]
    sheet.cell(row=1, column=(sheet.max_column or 0) + 1).value = header
    workbook.save(path)
    workbook.close()
    return path


def _outcomes(plan, sheet: str) -> list[str]:
    return [p.outcome for p in plan.rows if p.row.sheet == sheet and p.row.is_matter_row]


def test_a_shifted_sheet_plans_nothing_and_its_rows_wait_for_review(tmp_path: Path) -> None:
    plan = build_plan(_insert_a_column(_register(tmp_path)))

    shifted = [p for p in plan.rows if p.row.sheet == "2026" and p.row.is_matter_row]
    assert shifted
    assert {p.outcome for p in shifted} == {RowOutcome.REVIEW_REQUIRED.value}
    assert all(p.note == HEADER_DRIFT_NOTE for p in shifted)
    assert all(Anomaly.PARSE_ERROR.value in p.row.anomalies for p in shifted)
    assert any("2026" in finding for finding in plan.structure_findings)
    # The sheet beside it, whose header matches, still plans normally.
    assert _outcomes(plan, "2025") == [RowOutcome.WOULD_CREATE.value]


def test_an_uncontracted_column_after_the_contract_still_imports(tmp_path: Path) -> None:
    plan = build_plan(_append_a_column(_register(tmp_path)))

    assert set(_outcomes(plan, "2026")) == {RowOutcome.WOULD_CREATE.value}
    assert any("leping ei kirjelda" in finding for finding in plan.structure_findings)


def test_apply_refuses_a_shifted_sheet_and_never_guesses_its_rows(tmp_path: Path, settings) -> None:
    settings.REAL_DATA_ALLOWED = True
    path = _insert_a_column(_register(tmp_path))
    reports = str(tmp_path / "reports")

    with pytest.raises(CommandError, match="rows still need review"):
        call_command(
            "import_legacy_register",
            str(path),
            "--apply",
            "--report-dir",
            reports,
            stdout=io.StringIO(),
        )
    assert not Matter.all_objects.exists()

    # Accepting review rows imports the matching sheet and leaves the shifted one out
    # — with the plan digest the dry run printed (docs/adr/0184).
    from app.legacy_import.register_approval import register_plan_digest
    from app.legacy_import.resolution import MappingTables

    digest = register_plan_digest(build_plan(path), MappingTables.empty())
    call_command(
        "import_legacy_register",
        str(path),
        "--apply",
        "--accept-review-rows",
        "--expect-plan-sha256",
        digest,
        "--report-dir",
        reports,
        stdout=io.StringIO(),
    )
    assert sorted(Matter.all_objects.values_list("reference_year", flat=True)) == [2025]


def test_the_pilot_refuses_a_shifted_2026_sheet(tmp_path: Path) -> None:
    path = _insert_a_column(_register(tmp_path))

    with pytest.raises(PilotError, match="header does not match its contract"):
        read_pilot_sheet(path)


def test_the_pilot_still_reads_a_sheet_with_an_extra_trailing_column(tmp_path: Path) -> None:
    reading = read_pilot_sheet(_append_a_column(_register(tmp_path)))

    assert sorted(reading.rows) == ["2026_1", "2026_2"]


def test_the_delta_stops_rather_than_compare_shifted_columns(tmp_path: Path) -> None:
    path = _insert_a_column(_register(tmp_path))

    with pytest.raises(DeltaRefused, match="2026"):
        build_delta(workbook_path=path)
