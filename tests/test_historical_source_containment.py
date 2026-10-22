"""The historical readers stay inside the corpus they were given, and within budget.

Two findings of the 2026-10-22 security review, both about operator-supplied
historical sources (docs/adr/0178):

* **Containment** (csf_58b192a5e05041464608bf4f). ``captureId``,
  ``archivePageKey`` and ``originalPath`` are strings from the OneNote
  archive's own JSON, and a directory input is whatever the folder holds. A
  ``..`` in one of them, or a symbolic link planted inside the archive, used to
  read a file outside it and store it as evidence with the archive's
  provenance. Every read now goes through `safe_paths`.
* **Resource limits** (csf_38e223a165f6a7a5e91a3fa8). The opinions ZIP was
  inflated member by member with no bound on count, size or ratio, and the
  register workbook was parsed whole whatever its declared extent.

Everything is synthetic and built in ``tmp_path``. The bombs are small on disk:
megabytes of zeros deflate to kilobytes, and the limits are lowered through
settings so no test inflates more than a few MiB. Tests that need a symbolic
link or a FIFO skip only where the host cannot make one (Windows without the
symlink privilege); CI runs on Linux, where every one of them runs, and the
link tests run there against both the ``dir_fd`` walk and the ``lstat``
fallback.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import time
import unicodedata
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.legacy_import import parser as register_parser
from app.legacy_import import safe_paths
from app.legacy_import.limits import SourceLimitExceeded, declared_extent
from app.legacy_import.onenote_archive import ArchiveError, OneNoteArchive
from app.legacy_import.opinion_sources import (
    ArchiveReader,
    OpinionSourceError,
    read_kodadash_artifact,
    read_opinion_archive,
)
from app.legacy_import.parser import RegisterWorkbook
from app.legacy_import.safe_paths import UnsafeSourcePath
from tests import synthetic_historical
from tests import synthetic_opinions as syn

SECRET = b"outside the corpus: must never be read"

# ---------------------------------------------------------------------------
# What the host can make
# ---------------------------------------------------------------------------


def _can_symlink() -> bool:
    import tempfile

    with tempfile.TemporaryDirectory() as scratch:
        target = Path(scratch) / "t"
        target.write_bytes(b"")
        try:
            os.symlink(target, Path(scratch) / "l")
        except (OSError, NotImplementedError):
            return False
    return True


CAN_SYMLINK = _can_symlink()
needs_symlink = pytest.mark.skipif(
    not CAN_SYMLINK, reason="symlink creation needs a privilege on Windows"
)
needs_fifo = pytest.mark.skipif(not hasattr(os, "mkfifo"), reason="no FIFOs on this host")

#: Both ways of walking a path: the ``dir_fd`` one where the host has it, and
#: the ``lstat`` fallback (the only one on Windows) everywhere.
STRATEGIES = [
    pytest.param(
        True,
        id="dir_fd-walk",
        marks=pytest.mark.skipif(
            not safe_paths.STRONG_WALK, reason="no dir_fd/O_NOFOLLOW on this host"
        ),
    ),
    pytest.param(False, id="lstat-fallback"),
]


@pytest.fixture(params=STRATEGIES)
def walk(request, monkeypatch):
    monkeypatch.setattr(safe_paths, "STRONG_WALK", request.param)
    return request.param


@pytest.fixture
def outside(tmp_path) -> Path:
    """A directory beside the corpus holding the file nothing may read."""
    directory = tmp_path / "outside"
    directory.mkdir()
    (directory / "secret.pdf").write_bytes(SECRET)
    (directory / "original").mkdir()
    (directory / "original" / "ettepanek.pdf").write_bytes(SECRET)
    return directory


# ---------------------------------------------------------------------------
# The path rules
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "raw",
    [
        "",
        ".",
        "..",
        "../secret.pdf",
        "a/../../secret.pdf",
        "a/./b.pdf",
        "a//b.pdf",
        "a/",
        "/etc/passwd",
        "C:/Windows/win.ini",
        "C:secret.pdf",
        "\\\\server\\share\\x.pdf",
        "resources\\..\\x.pdf",
        "a/b\x00.pdf",
        "a/b\n.pdf",
        "a/b\x7f.pdf",
        "a/trailing-dot.",
        "a/trailing-space ",
        "a/stream.pdf:hidden",
        "\uff0e\uff0e/secret.pdf",  # fullwidth dots fold to ".."
        "a\uff0fb.pdf",  # a fullwidth solidus folds to "/"
        None,
        42,
    ],
)
def test_an_unsafe_relative_path_is_refused(raw):
    with pytest.raises(UnsafeSourcePath, match="Refusing an unsafe source path"):
        safe_paths.relative_parts(raw)


@pytest.mark.parametrize(
    "raw",
    [
        "resources/r1/original/Ettepanek.pdf",
        "Opinions/2021-02-01 - Saaja - Pealkiri.pdf",
        unicodedata.normalize("NFC", "Õigusloome/Määruse eelnõu ülevõtmise kohta .pdf"),
        unicodedata.normalize("NFD", "Õigusloome/Määruse eelnõu ülevõtmise kohta .pdf"),
        "Opinions/pr\x81gila ja kaevandamisj\u201e\u201etme.pdf",  # the real archive's mojibake
        "..hidden/.x/a..b.pdf",
    ],
)
def test_a_legitimate_name_passes_exactly_as_written(raw):
    assert "/".join(safe_paths.relative_parts(raw)) == raw


# ---------------------------------------------------------------------------
# The OneNote archive
# ---------------------------------------------------------------------------


@pytest.fixture
def onenote(tmp_path) -> Path:
    return synthetic_historical.build_archive(tmp_path / "corpus" / "onenote-desktop-archive")


def _rewrite_json(path: Path, **changes) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    payload.update(changes)
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


@pytest.fixture
def opened(monkeypatch) -> list[str]:
    """Every path `safe_paths` opens, so a test can prove the secret never was."""
    seen: list[str] = []
    real = safe_paths.open_file

    def recording(root, relative):
        seen.append(str(Path(root) / relative))
        return real(root, relative)

    monkeypatch.setattr(safe_paths, "open_file", recording)
    return seen


@pytest.mark.parametrize(
    "original_path",
    [
        "resources\\..\\..\\..\\..\\outside\\secret.pdf",
        "../../../../outside/secret.pdf",
        "C:\\outside\\secret.pdf",
        "/outside/secret.pdf",
        "resources/r-exact-1/original/./ettepanek.pdf",
    ],
)
def test_an_unsafe_original_path_is_refused_when_the_page_is_read(
    onenote, outside, opened, original_path
):
    _rewrite_json(
        onenote / "pages/p-exact/resources/r-exact-1/resource.json", originalPath=original_path
    )
    archive = OneNoteArchive(onenote)

    with pytest.raises(ArchiveError, match="originalPath: Refusing an unsafe source path"):
        archive.page("p-exact")
    assert not any("outside" in path for path in opened)


def test_an_unsafe_capture_id_is_refused_before_anything_is_read(onenote, outside, opened):
    _rewrite_json(onenote / "pages/p-exact/current.json", captureId="../../../../outside")
    with pytest.raises(ArchiveError, match="captureId: Refusing an unsafe source path"):
        OneNoteArchive(onenote).page("p-exact")
    assert not any("outside" in path for path in opened)


def test_an_unsafe_archive_page_key_is_refused(onenote):
    """Every later read of the page is joined to this key, so it must be one segment."""
    _rewrite_json(onenote / "pages/p-exact/page.json", archivePageKey="../p-shared")
    with pytest.raises(ArchiveError, match="archivePageKey: Refusing an unsafe source path"):
        OneNoteArchive(onenote).page("p-exact")


@pytest.mark.parametrize("page_key", ["../p-exact", "p-exact/../p-shared", "", "/abs"])
def test_an_unsafe_page_key_is_refused(onenote, page_key):
    with pytest.raises(ArchiveError, match="page key: Refusing an unsafe source path"):
        OneNoteArchive(onenote).page(page_key)


def test_a_stored_path_is_checked_again_when_it_is_read(onenote, outside, opened):
    """Materialisation reads by the path the catalogue stored, not by a fresh page read."""
    archive = OneNoteArchive(onenote)
    with pytest.raises(ArchiveError, match="resource path: Refusing an unsafe source path"):
        archive.read_resource_at("p-exact", "../../../../outside/secret.pdf")
    with pytest.raises(ArchiveError, match="page key: Refusing an unsafe source path"):
        archive.read_resource_at("..", "outside/secret.pdf")
    assert opened == []


def _symlink(link: Path, target: Path) -> None:
    if link.is_dir() and not link.is_symlink():
        shutil.rmtree(link)
    elif link.exists() or link.is_symlink():
        link.unlink()
    os.symlink(target, link, target_is_directory=target.is_dir())


@needs_symlink
def test_a_resource_that_is_a_link_out_of_the_archive_is_refused(walk, onenote, outside):
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    resource = next(r for r in page.resources if r.resource_key == "r-exact-1")
    _symlink(onenote / "pages/p-exact" / resource.relative_path, outside / "secret.pdf")

    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        archive.read_resource("p-exact", resource)


@needs_symlink
def test_a_directory_on_the_way_that_is_a_link_is_refused(walk, onenote, outside):
    """The intermediate case: ``original/`` swapped for a link to a look-alike folder."""
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    resource = next(r for r in page.resources if r.resource_key == "r-exact-1")
    _symlink(onenote / "pages/p-exact/resources/r-exact-1/original", outside / "original")

    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        archive.read_resource("p-exact", resource)


@needs_symlink
def test_a_linked_page_record_or_capture_is_refused(walk, onenote, outside):
    (outside / "page.json").write_text(json.dumps({"title": "planted"}), encoding="utf-8")
    _symlink(onenote / "pages/p-exact/page.json", outside / "page.json")
    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        OneNoteArchive(onenote).page("p-exact")

    capture = onenote / "pages/p-shared/captures/c0000001"
    shutil.copytree(capture, outside / "capture")
    _symlink(capture, outside / "capture")
    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        OneNoteArchive(onenote).page("p-shared")


@needs_symlink
def test_a_linked_page_directory_is_refused_not_followed(walk, onenote, outside):
    shutil.copytree(onenote / "pages/p-exact", outside / "p-planted")
    os.symlink(outside / "p-planted", onenote / "pages/p-planted", target_is_directory=True)
    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        OneNoteArchive(onenote).page("p-planted")


@needs_fifo
def test_a_fifo_in_place_of_a_resource_is_refused_without_blocking(walk, onenote):
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    resource = next(r for r in page.resources if r.resource_key == "r-exact-1")
    target = onenote / "pages/p-exact" / resource.relative_path
    target.unlink()
    os.mkfifo(target)

    with pytest.raises(ArchiveError, match="not a regular file"):
        archive.read_resource("p-exact", resource)


# -- the window between the check and the open ------------------------------


def _swap_after_check(monkeypatch, swap) -> None:
    """Run `swap` after the path is checked and before it is opened."""
    for name in ("_walk_strong", "_walk_fallback"):
        real = getattr(safe_paths, name)

        def checked_then_swapped(root, parts, raw, _real=real):
            result = _real(root, parts, raw)
            swap()
            return result

        monkeypatch.setattr(safe_paths, name, checked_then_swapped)


def _resource_file(onenote: Path) -> tuple[OneNoteArchive, object, Path]:
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    resource = next(r for r in page.resources if r.resource_key == "r-exact-1")
    return archive, resource, onenote / "pages/p-exact" / resource.relative_path


@needs_symlink
def test_a_file_swapped_for_a_link_after_the_check_is_not_followed(
    walk, onenote, outside, monkeypatch
):
    archive, resource, target = _resource_file(onenote)

    def swap():
        target.rename(target.with_name("moved-away.pdf"))
        os.symlink(outside / "secret.pdf", target)

    _swap_after_check(monkeypatch, swap)
    with pytest.raises(ArchiveError, match="changed between the check and the open"):
        archive.read_resource("p-exact", resource)


@needs_symlink
def test_a_directory_swapped_for_a_link_after_the_check_cannot_redirect_the_read(
    walk, onenote, outside, monkeypatch
):
    """Either the read is refused, or it returns the file that was checked — never the secret.

    The ``dir_fd`` walk already holds the real directory, so it reads the
    original bytes; the fallback sees a different file behind the path and
    refuses.
    """
    archive, resource, target = _resource_file(onenote)
    original_bytes = target.read_bytes()
    directory = target.parent

    def swap():
        directory.rename(directory.with_name("original-moved"))
        os.symlink(outside / "original", directory, target_is_directory=True)

    _swap_after_check(monkeypatch, swap)
    try:
        content = archive.read_resource("p-exact", resource)
    except ArchiveError as error:
        assert "changed between the check and the open" in str(error)
    else:
        assert content == original_bytes
        assert content != SECRET


def test_a_file_replaced_after_the_check_is_refused(walk, onenote, monkeypatch):
    """Runs on every host: the identity comparison, without needing a link."""
    archive, resource, target = _resource_file(onenote)

    def swap():
        target.rename(target.with_name("moved-away.pdf"))
        target.write_bytes(SECRET)

    _swap_after_check(monkeypatch, swap)
    with pytest.raises(ArchiveError, match="changed between the check and the open"):
        archive.read_resource("p-exact", resource)


# -- the page record the evidence pilot reads ---------------------------------


def test_the_page_record_is_read_through_the_archive(onenote, opened):
    """`evidence_pilot` takes `page.json` whole from here, by the archive's own key."""
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    record = archive.page_record(page.page_key)
    assert record["archivePageKey"] == "p-exact"
    assert any(path.endswith("page.json") for path in opened)
    with pytest.raises(ArchiveError, match="page key: Refusing an unsafe source path"):
        archive.page_record("../p-exact")


@needs_symlink
def test_a_page_record_swapped_for_a_link_after_the_page_was_read_is_refused(
    walk, onenote, outside, opened
):
    """The second read of ``page.json`` is checked again, not trusted from the first."""
    (outside / "page.json").write_text(json.dumps({"title": "planted"}), encoding="utf-8")
    archive = OneNoteArchive(onenote)
    page = archive.page("p-exact")
    _symlink(onenote / "pages/p-exact/page.json", outside / "page.json")
    opened.clear()

    with pytest.raises(ArchiveError, match="a symbolic link inside the source"):
        archive.page_record(page.page_key)
    assert opened == []


def test_no_legacy_reader_opens_an_archive_page_by_a_bare_path():
    """Every read under the OneNote pages goes through `safe_paths` (docs/adr/0178).

    A source guard, because the defect it pins is a single line: joining the
    archive's own `archivePageKey` to `pages_root` and calling ``read_text`` was
    how `evidence_pilot` read ``page.json`` until 2026-10-22.
    """
    import re

    from django.conf import settings

    reader = re.compile(r"(read_text|read_bytes|open)\(")
    offenders = []
    for path in sorted((Path(settings.BASE_DIR) / "app" / "legacy_import").rglob("*.py")):
        if path.name in {"safe_paths.py", "limits.py"}:
            continue
        for number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
            if "pages_root" in line and reader.search(line):
                offenders.append(f"{path.name}:{number}: {line.strip()}")
    assert offenders == []


# -- legitimate archives still read -------------------------------------------


@pytest.mark.parametrize("form", ["NFC", "NFD"])
def test_an_estonian_filename_still_reads_in_either_normal_form(tmp_path, form):
    name = unicodedata.normalize(form, "Märkused õigusloome ülevõtmise kohta.pdf")
    page = synthetic_historical.page_specs()[0]
    page.resources = [
        synthetic_historical.ResourceSpec(
            resource_key="r-uni", filename=name, content=b"%PDF-1.4 uni", ordinal=3
        )
    ]
    archive = OneNoteArchive(synthetic_historical.build_archive(tmp_path / "a", [page]))

    read = archive.page(page.page_key)
    (resource,) = read.resources
    assert resource.relative_path == f"resources/r-uni/original/{name}"
    assert archive.read_resource(read.page_key, resource) == b"%PDF-1.4 uni"
    assert archive.read_page_xml(read).startswith(b"<?xml")


# ---------------------------------------------------------------------------
# The opinions archive as a directory
# ---------------------------------------------------------------------------


@pytest.fixture
def opinion_dir(tmp_path) -> Path:
    root = tmp_path / "corpus" / "Opinions"
    root.mkdir(parents=True)
    for date, recipient, title in [
        ("2021-02-01", "Rahandusministeerium", "Maksumuudatused"),
        ("2022-03-04", "MKM", "Õigusloome ülevõtmine"),
    ]:
        (root / f"{date} - {recipient} - {title}.pdf").write_bytes(syn.pdf_bytes(title))
    return root


@needs_symlink
def test_a_link_in_a_directory_archive_is_refused_not_catalogued(walk, opinion_dir, outside):
    os.symlink(outside / "secret.pdf", opinion_dir / "2023-01-01 - X - Planted.pdf")
    with pytest.raises(OpinionSourceError, match="a symbolic link inside the source"):
        read_opinion_archive(opinion_dir)


@needs_symlink
def test_a_linked_subdirectory_in_a_directory_archive_is_refused(walk, opinion_dir, outside):
    os.symlink(outside / "original", opinion_dir / "Lisad", target_is_directory=True)
    with pytest.raises(OpinionSourceError, match="a symbolic link inside the source"):
        read_opinion_archive(opinion_dir)


@needs_symlink
def test_the_directory_reader_refuses_a_link_and_a_traversal(walk, opinion_dir, outside):
    reader = ArchiveReader(opinion_dir)
    with pytest.raises(OpinionSourceError, match="Refusing an unsafe source path"):
        reader.read("../outside/secret.pdf")
    os.symlink(outside / "secret.pdf", opinion_dir / "linked.pdf")
    with pytest.raises(OpinionSourceError, match="a symbolic link inside the source"):
        reader.read("linked.pdf")


def test_the_directory_reader_refuses_a_traversal_on_every_host(opinion_dir, outside):
    with pytest.raises(OpinionSourceError, match="Refusing an unsafe source path"):
        ArchiveReader(opinion_dir).read("../../outside/secret.pdf")


@needs_fifo
def test_a_fifo_in_a_directory_archive_is_refused(walk, opinion_dir):
    os.mkfifo(opinion_dir / "pipe.pdf")
    with pytest.raises(OpinionSourceError, match="not a regular file"):
        read_opinion_archive(opinion_dir)


def test_a_legitimate_directory_archive_reads_unchanged(walk, opinion_dir):
    """Same occurrences, same manifest identity, same bytes back — Unicode names included."""
    nfd = unicodedata.normalize("NFD", "2023-05-06 - Kliimaministeerium - Jäätmete äöü.pdf")
    (opinion_dir / nfd).write_bytes(syn.pdf_bytes("nfd"))

    digest, occurrences = read_opinion_archive(opinion_dir)

    expected = sorted(p for p in opinion_dir.rglob("*") if p.is_file())
    assert [o.relative_path for o in occurrences] == [
        p.relative_to(opinion_dir).as_posix() for p in expected
    ]
    manifest = "\n".join(f"{o.relative_path}\t{o.sha256}" for o in occurrences)
    assert digest == hashlib.sha256(manifest.encode("utf-8")).hexdigest()
    reader = ArchiveReader(opinion_dir)
    for occurrence, path in zip(occurrences, expected, strict=True):
        assert occurrence.sha256 == hashlib.sha256(path.read_bytes()).hexdigest()
        assert occurrence.size_bytes == path.stat().st_size
        assert occurrence.detected_type == "application/pdf"
        assert reader.read(occurrence.relative_path) == path.read_bytes()
    assert reader.read("absent.pdf") is None


def test_a_directory_archive_past_its_member_limit_is_refused(opinion_dir, settings):
    settings.LEGACY_IMPORT_MAX_ARCHIVE_MEMBERS = 1
    with pytest.raises(OpinionSourceError, match="the limit is 1"):
        read_opinion_archive(opinion_dir)


# ---------------------------------------------------------------------------
# The opinions ZIP
# ---------------------------------------------------------------------------

MIB = 1024 * 1024


def _deflated(path: Path, entries: list[tuple[str, bytes]]) -> Path:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return path


@pytest.fixture
def forbid_inflation(monkeypatch):
    """Call once the archive is written: from then on, opening any member fails the test."""

    def refuse(*args, **kwargs):
        raise AssertionError("a member was inflated before the archive was refused")

    def forbid() -> None:
        monkeypatch.setattr(zipfile.ZipFile, "open", refuse)

    return forbid


def test_a_zip_with_too_many_members_is_refused_before_any_is_read(
    tmp_path, settings, forbid_inflation
):
    settings.LEGACY_IMPORT_MAX_ARCHIVE_MEMBERS = 5
    path = _deflated(
        tmp_path / "many.zip", [(f"Opinions/{n}.pdf", syn.pdf_bytes(str(n))) for n in range(6)]
    )
    forbid_inflation()
    with pytest.raises(OpinionSourceError, match="has 6 members; the limit is 5"):
        read_opinion_archive(path)
    with pytest.raises(OpinionSourceError, match="has 6 members; the limit is 5"):
        ArchiveReader(path)


def test_a_zip_bomb_by_ratio_is_refused_before_any_is_read(tmp_path, forbid_inflation):
    """2 MiB of zeros deflates to about 2 KB: a ratio near 1000, five times the default."""
    path = _deflated(tmp_path / "ratio.zip", [("Opinions/bomb.pdf", bytes(2 * MIB))])
    assert path.stat().st_size < 64 * 1024
    forbid_inflation()
    with pytest.raises(OpinionSourceError, match="implausible compression ratio"):
        read_opinion_archive(path)


def test_a_zip_whose_members_inflate_past_the_total_is_refused_before_any_is_read(
    tmp_path, settings, forbid_inflation
):
    settings.LEGACY_IMPORT_MAX_COMPRESSION_RATIO = 100_000
    settings.LEGACY_IMPORT_MAX_TOTAL_UNCOMPRESSED_BYTES = 3 * MIB
    path = _deflated(
        tmp_path / "volume.zip",
        [("Opinions/a.pdf", bytes(2 * MIB)), ("Opinions/b.pdf", bytes(2 * MIB))],
    )
    forbid_inflation()
    with pytest.raises(OpinionSourceError, match="declares more than 3145728 bytes"):
        read_opinion_archive(path)


def test_one_member_past_its_limit_is_refused_before_any_is_read(
    tmp_path, settings, forbid_inflation
):
    settings.LEGACY_IMPORT_MAX_COMPRESSION_RATIO = 100_000
    settings.LEGACY_IMPORT_MAX_MEMBER_BYTES = 1 * MIB
    path = _deflated(tmp_path / "member.zip", [("Opinions/a.pdf", bytes(2 * MIB))])
    forbid_inflation()
    with pytest.raises(OpinionSourceError, match="declares 2097152 bytes; the limit is 1048576"):
        read_opinion_archive(path)


def test_an_archive_file_past_its_size_limit_is_refused_unopened(tmp_path, settings):
    settings.LEGACY_IMPORT_MAX_ARCHIVE_BYTES = 100
    path = _deflated(tmp_path / "big.zip", [("Opinions/a.pdf", syn.pdf_bytes("a" * 200))])
    with pytest.raises(OpinionSourceError, match=r"the archive is \d+ bytes; the limit is 100"):
        read_opinion_archive(path)


def _understate_size(path: Path, declared: int) -> None:
    """Rewrite one member's uncompressed size in its local header and directory record."""
    with zipfile.ZipFile(path) as archive:
        (info,) = archive.infolist()
        local, directory = info.header_offset, archive.start_dir
    raw = bytearray(path.read_bytes())
    assert (
        raw[local : local + 4] == b"PK\x03\x04" and raw[directory : directory + 4] == b"PK\x01\x02"
    )
    raw[local + 22 : local + 26] = declared.to_bytes(4, "little")
    raw[directory + 24 : directory + 28] = declared.to_bytes(4, "little")
    path.write_bytes(bytes(raw))


def test_a_member_that_inflates_past_its_declared_size_is_refused_while_streaming(tmp_path):
    """The declared size is a claim: the bytes actually produced are what is counted."""
    path = _deflated(tmp_path / "liar.zip", [("Opinions/a.pdf", bytes(512 * 1024))])
    _understate_size(path, 1000)
    with pytest.raises(OpinionSourceError, match=r"(damaged|inflates past)"):
        read_opinion_archive(path)

    with zipfile.ZipFile(tmp_path / "liar.zip") as archive:
        assert archive.infolist()[0].file_size == 1000
    reader = ArchiveReader(path)
    with pytest.raises(OpinionSourceError, match=r"(damaged|inflates past)"):
        reader.read("Opinions/a.pdf")


def test_a_legitimate_zip_reads_exactly_as_before(tmp_path):
    opinions = [
        syn.opinion(date="2021-02-01", recipient="Rahandusministeerium", title="Maksumuudatused"),
        syn.opinion(date="2022-03-04", recipient="MKM", title="Õigusloome ülevõtmine"),
    ]
    path = syn.write_archive(tmp_path / "ok.zip", opinions)
    digest, occurrences = read_opinion_archive(path)

    assert digest == hashlib.sha256(path.read_bytes()).hexdigest()
    reader = ArchiveReader(path)
    with zipfile.ZipFile(path) as archive:
        for occurrence in occurrences:
            data = archive.read(occurrence.relative_path)
            assert occurrence.sha256 == hashlib.sha256(data).hexdigest()
            assert occurrence.size_bytes == len(data)
            assert reader.read(occurrence.relative_path) == data


# ---------------------------------------------------------------------------
# Workbooks
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("ref", "extent"),
    [("A1:WVP313", (313, 16136)), ("A1", (1, 1)), ("$B$2:$C$5", (4, 2)), ("bogus", None)],
)
def test_a_declared_range_is_read_as_rows_and_columns(ref, extent):
    assert declared_extent(ref) == extent


def _far_cell_workbook(path: Path) -> Path:
    """Two cells, one of them at XFD1048576: a few KB on disk, 17 billion cells of range."""
    book = Workbook()
    sheet = book.active
    sheet.title = "2026"
    sheet["A1"] = "Viide"
    sheet.cell(row=1_048_576, column=16_384, value="far")
    book.save(path)
    return path


def _rewrite_part(path: Path, part: str, old: bytes, new: bytes) -> None:
    with zipfile.ZipFile(path) as source:
        entries = [(info, source.read(info)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for info, data in entries:
            if info.filename == part:
                assert old in data
                data = data.replace(old, new)
            target.writestr(info, data)


def test_an_extreme_declared_range_is_refused_before_the_workbook_is_parsed(tmp_path, monkeypatch):
    path = _far_cell_workbook(tmp_path / "far.xlsx")
    monkeypatch.setattr(
        register_parser,
        "load_workbook",
        lambda *a, **k: (_ for _ in ()).throw(AssertionError("parsed before the check")),
    )
    started = time.monotonic()
    with pytest.raises(SourceLimitExceeded, match="spans 1048576 rows; the limit is 200000"):
        RegisterWorkbook(path)
    assert time.monotonic() - started < 5


def test_a_range_hidden_from_the_dimension_is_caught_once_loaded(tmp_path):
    """A workbook that declares ``A1`` but holds a cell at XFD1048576.

    The pre-check cannot see it; the loaded extent can, and is refused before
    any row loop asks openpyxl to create the cells in between.
    """
    path = _far_cell_workbook(tmp_path / "hidden.xlsx")
    _rewrite_part(path, "xl/worksheets/sheet1.xml", b'ref="A1:XFD1048576"', b'ref="A1"')
    started = time.monotonic()
    with pytest.raises(SourceLimitExceeded, match="spans 1048576 rows"):
        RegisterWorkbook(path)
    assert time.monotonic() - started < 10


def test_a_workbook_with_too_many_sheets_is_refused(tmp_path, settings):
    settings.LEGACY_IMPORT_MAX_WORKSHEETS = 3
    book = Workbook()
    for year in (2023, 2024, 2025):
        book.create_sheet(str(year))
    book.save(tmp_path / "sheets.xlsx")
    with pytest.raises(SourceLimitExceeded, match="has 4 worksheets; the limit is 3"):
        RegisterWorkbook(tmp_path / "sheets.xlsx")


def test_a_workbook_bomb_is_refused_before_it_is_parsed(tmp_path, settings):
    settings.LEGACY_IMPORT_MAX_WORKBOOK_UNCOMPRESSED_BYTES = 1 * MIB
    book = Workbook()
    for row in range(1, 80):
        # Distinct strings, so the shared-string table cannot fold them into one.
        book.active.cell(row=row, column=1, value=f"{row:05d}" + "x" * 32_000)
    book.save(tmp_path / "bomb.xlsx")
    with pytest.raises(SourceLimitExceeded, match="workbook"):
        RegisterWorkbook(tmp_path / "bomb.xlsx")


def _kodadash_with(tmp_path: Path, rows: int) -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "02_source_binding_audit"
    sheet.append(["content_id", "file_sha256"])
    for n in range(rows):
        sheet.append([f"c{n}", hashlib.sha256(str(n).encode()).hexdigest()])
    path = tmp_path / "kodadash.xlsx"
    book.save(path)
    return path


def test_the_kodadash_workbook_has_the_same_bounds(tmp_path, settings):
    path = _far_cell_workbook(tmp_path / "far.xlsx")
    with pytest.raises(OpinionSourceError, match="spans 1048576 rows"):
        read_kodadash_artifact(path)

    settings.LEGACY_IMPORT_MAX_SHEET_ROWS = 5
    path = _kodadash_with(tmp_path, rows=10)
    with pytest.raises(OpinionSourceError, match="spans 11 rows; the limit is 5"):
        read_kodadash_artifact(path)


def test_kodadash_rows_are_counted_as_they_stream_when_no_range_is_declared(tmp_path, settings):
    path = _kodadash_with(tmp_path, rows=10)
    _rewrite_part(path, "xl/worksheets/sheet1.xml", b'<dimension ref="A1:B11"/>', b"")
    assert len(read_kodadash_artifact(path)[1]) == 10

    settings.LEGACY_IMPORT_MAX_SHEET_ROWS = 5
    with pytest.raises(OpinionSourceError, match="more than 5 rows"):
        read_kodadash_artifact(path)


# ---------------------------------------------------------------------------
# Materialisation (database)
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_materialising_a_tampered_stored_path_fails_that_file_and_reads_nothing_outside(
    tmp_path, outside
):
    """`archive_relative_path` is read back from the database at materialisation.

    A row that names a path out of the archive fails as one file — the run goes
    on — and no document is ever made from the bytes it pointed at.
    """
    from app.documents.models import DocumentVersion
    from app.legacy_import.historical_apply import (
        apply_structure,
        materialise_resources,
        open_batch,
    )
    from app.legacy_import.historical_plan import build_plan
    from app.legacy_import.source_pages import (
        LegacySourceResource,
        LegacySourceResourceImport,
        ResourceImportState,
    )
    from app.matters.enums import MatterOrigin, RecordMode
    from tests import factories

    corpus = synthetic_historical.build_corpus(tmp_path / "source")
    for reference, title, _ in synthetic_historical.EXACT_MATCHES:
        if reference != "2019_9":
            factories.MatterFactory(
                title=title,
                reference_year=2019,
                reference_number=int(reference.split("_")[1]),
                record_mode=RecordMode.ARCHIVE,
                origin=MatterOrigin.LEGACY_IMPORT,
            )
    plan = build_plan(
        excel_path=corpus["excel_path"],
        archive_root=corpus["archive_root"],
        audit_root=corpus["audit_root"],
    )
    archive = OneNoteArchive(corpus["archive_root"])
    batch = open_batch(plan, importer_version="test")
    apply_structure(plan, batch=batch, archive=archive)
    # Deep enough to leave tmp_path whatever the corpus layout.
    LegacySourceResource.objects.filter(resource_key="r-exact-1").update(
        archive_relative_path="../" * 12 + "outside/secret.pdf"
    )

    report = materialise_resources(archive=archive, batch=batch)

    failed = LegacySourceResourceImport.objects.get(resource__resource_key="r-exact-1")
    assert failed.state == ResourceImportState.FAILED
    assert failed.error_code == "ArchiveError"
    assert "Refusing an unsafe source path" in failed.error_detail
    assert len(report.failures) == 1
    assert report.documents_created >= 5
    secret_sha = hashlib.sha256(SECRET).hexdigest()
    assert not DocumentVersion.objects.filter(sha256=secret_sha).exists()
