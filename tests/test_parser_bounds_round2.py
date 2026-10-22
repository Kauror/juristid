"""The remaining document parsers stay bounded on hostile input (docs/adr/0182).

Four findings of the second 2026-10-22 scan, each held to two things: the
hostile file is refused *before* the allocation it was built to cause, and an
ordinary file reads exactly as it did — the old algorithms are kept here as the
definition of "exactly":

* csf_9b7b703ca586c216499d33bd — an image over the pixel ceiling is refused
  from its header, before any decode, on every image path;
* csf_5fb4e796365c901966387382 — an Office package's directory is checked from
  its end records before `zipfile` builds an entry for any of it;
* csf_ea514a0f447d737dfb1bdfb9 — plain text is decoded and split in bounded
  pieces, never whole;
* csf_eb1ed7815b093aba8872620b — DOCX parts are budgeted before python-docx
  parses them, and paragraphs, rows and cells are charged as they are walked.

Every hostile file is generated in the test, small; limits are lowered to meet
them rather than files raised to meet the limits.
"""

from __future__ import annotations

import codecs
import io
import random
import struct
import sys
import tracemalloc
import zipfile

import pytest

from app.documents.extraction import office as office_module
from app.documents.extraction import text as text_module
from app.documents.extraction.base import SourceFile, normalise, registry
from app.documents.extraction.errors import ExtractionFailed
from app.documents.extraction.limits import current_limits, preflight_zip_directory

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT = "text/plain"
PNG = "image/png"


def parse(mime_type: str, content: bytes):
    parser = registry.for_mime_type(mime_type)
    assert parser is not None
    return parser.parse(SourceFile(content=content, filename="fail", mime_type=mime_type))


def refused_with(code: str, call, *args) -> ExtractionFailed:
    with pytest.raises(ExtractionFailed) as error:
        call(*args)
    assert error.value.code == code, (error.value.code, error.value.detail)
    return error.value


def outcome(call, *args):
    """What a parse produced, as data: its fragments and metadata, or its refusal."""
    try:
        result = call(*args)
    except ExtractionFailed as error:
        return ("refused", error.code)
    except Exception as error:  # a crash is an outcome to compare too
        return ("crashed", type(error).__name__)
    return (
        "read",
        [
            [(f.text, f.locator, f.locator_label, f.locator_kind) for f in payload.fragments]
            for payload in result.derivatives
        ],
        [payload.metadata for payload in result.derivatives],
    )


# ---------------------------------------------------------------------------
# Images — csf_9b7b703ca586c216499d33bd
# ---------------------------------------------------------------------------


def _png(width: int, height: int) -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (width, height), "white").save(buffer, format="PNG")
    return buffer.getvalue()


@pytest.fixture
def no_decode(monkeypatch):
    """Fail loudly if Pillow starts decoding pixels."""
    from PIL import ImageFile

    def decoded(self, *args, **kwargs):
        raise AssertionError("the image was decoded")

    monkeypatch.setattr(ImageFile.ImageFile, "load", decoded)


@pytest.fixture
def ocr_engine(monkeypatch):
    """OCR reported available, so the parser reaches its OCR path."""
    from app.documents.extraction import ocr

    monkeypatch.setattr(text_module, "ocr_is_available", lambda: True)
    monkeypatch.setattr(text_module, "ocr_engine_version", lambda: "tesseract test")
    monkeypatch.setattr(ocr, "ocr_engine_version", lambda: "tesseract test")


def test_an_image_in_pillows_warning_band_is_refused_before_it_is_decoded(
    settings, no_decode, ocr_engine
):
    """1.5 × the ceiling: Pillow only warned here and then decoded it."""
    settings.EXTRACTION_MAX_IMAGE_PIXELS = 10_000
    refused_with("image_too_large", parse, PNG, _png(150, 100))


def test_the_thumbnail_and_ocr_paths_each_refuse_from_the_header(settings, no_decode, ocr_engine):
    from app.documents.extraction.ocr import recognise_image
    from app.documents.extraction.thumbnails import image_thumbnail

    settings.EXTRACTION_MAX_IMAGE_PIXELS = 10_000
    content = _png(150, 100)
    limits = current_limits()

    assert image_thumbnail(content, limits) is None
    refused_with("image_too_large", recognise_image, content, limits)


def test_pillows_own_hard_error_now_fires_at_the_ceiling(settings):
    from PIL import Image

    from app.documents.extraction.limits import configure_image_safety

    settings.EXTRACTION_MAX_IMAGE_PIXELS = 10_000
    configure_image_safety(current_limits())
    try:
        with pytest.raises(Image.DecompressionBombError):
            Image.open(io.BytesIO(_png(150, 100)))
    finally:
        settings.EXTRACTION_MAX_IMAGE_PIXELS = 80_000_000
        configure_image_safety(current_limits())


def test_an_image_within_the_ceiling_still_gets_its_thumbnail(settings):
    from app.documents.extraction.thumbnails import image_thumbnail

    settings.EXTRACTION_MAX_IMAGE_PIXELS = 10_000
    thumbnail = image_thumbnail(_png(100, 100), current_limits())

    assert thumbnail is not None
    assert thumbnail.metadata == {"width": 100, "height": 100, "format": "PNG"}


# ---------------------------------------------------------------------------
# ZIP directories — csf_5fb4e796365c901966387382
# ---------------------------------------------------------------------------


def _zip(members: int, *, comment: bytes = b"", prefix: bytes = b"") -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_STORED) as archive:
        for index in range(members):
            archive.writestr(f"m{index}", b"")
        archive.comment = comment
    return prefix + buffer.getvalue()


@pytest.fixture
def no_zipfile(monkeypatch):
    """Fail loudly if `zipfile` starts reading a central directory."""

    def read(self):
        raise AssertionError("zipfile read the central directory")

    monkeypatch.setattr(zipfile.ZipFile, "_RealGetContents", read)


def test_many_tiny_members_are_refused_before_zipfile_reads_the_directory(settings, no_zipfile):
    settings.EXTRACTION_MAX_ARCHIVE_MEMBERS = 50
    refused_with("too_many_members", parse, DOCX, _zip(500))


def test_a_directory_that_understates_its_entry_count_is_still_counted(settings, no_zipfile):
    """The end record says one entry; the directory holds five hundred."""
    settings.EXTRACTION_MAX_ARCHIVE_MEMBERS = 50
    content = bytearray(_zip(500))
    end = content.rfind(b"PK\x05\x06")
    struct.pack_into("<2H", content, end + 8, 1, 1)
    refused_with("too_many_members", parse, DOCX, bytes(content))


def test_a_directory_larger_than_its_budget_is_refused_unread(settings, no_zipfile):
    settings.EXTRACTION_MAX_ZIP_DIRECTORY_BYTES = 1024
    refused_with("zip_directory_limit", parse, DOCX, _zip(40))


def test_a_zip64_sentinel_without_a_zip64_record_is_refused(no_zipfile):
    content = bytearray(_zip(3))
    end = content.rfind(b"PK\x05\x06")
    struct.pack_into("<L", content, end + 16, 0xFFFFFFFF)
    refused_with("invalid_container", parse, DOCX, bytes(content))


def _as_zip64(content: bytes, *, offset_shift: int = 0) -> bytes:
    """The same archive with a ZIP64 end record and locator, as large writers emit."""
    end = content.rfind(b"PK\x05\x06")
    (_sig, _disk, _start, here, total, size, offset, _comment) = struct.unpack_from(
        "<4s4H2LH", content, end
    )
    record_at = end
    record = struct.pack(
        "<4sQ2H2L4Q", b"PK\x06\x06", 44, 45, 45, 0, 0, here, total, size, offset + offset_shift
    )
    locator = struct.pack("<4sLQL", b"PK\x06\x07", 0, record_at, 1)
    eocd = struct.pack("<4s4H2LH", b"PK\x05\x06", 0, 0, 0xFFFF, 0xFFFF, 0xFFFFFFFF, 0xFFFFFFFF, 0)
    return content[:end] + record + locator + eocd


def test_a_consistent_zip64_archive_passes_exactly_as_zipfile_reads_it():
    content = _as_zip64(_zip(3))
    preflight_zip_directory(io.BytesIO(content), current_limits())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert len(archive.infolist()) == 3


def test_an_inconsistent_zip64_record_is_refused(no_zipfile):
    refused_with("invalid_container", parse, DOCX, _as_zip64(_zip(3), offset_shift=7))


@pytest.mark.parametrize(
    "content",
    [_zip(3), _zip(3, comment=b"kommentaar"), _zip(3, prefix=b"prepended bytes")],
    ids=["plain", "comment", "prepended"],
)
def test_archives_zipfile_reads_pass_the_preflight(content):
    preflight_zip_directory(io.BytesIO(content), current_limits())
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert len(archive.infolist()) == 3


def test_real_office_files_pass_the_preflight():
    from tests import synthetic_corpus as corpus

    for content in (corpus.draft_docx(), corpus.annex_xlsx(), corpus.briefing_pptx()):
        preflight_zip_directory(io.BytesIO(content), current_limits())


# ---------------------------------------------------------------------------
# Plain text — csf_ea514a0f447d737dfb1bdfb9
# ---------------------------------------------------------------------------


def reference_text_parse(content: bytes):
    """The 9a2e834a TextParser, whole-file decode and ``splitlines`` included."""
    from app.documents.enums import LocatorKind
    from app.documents.extraction.base import DerivativePayload, Fragment, ParseResult
    from app.documents.extraction.limits import guard_character_budget

    limits = current_limits()
    text, encoding = text_module.decode_text(content)
    lines = text.splitlines()
    fragments = []
    total = 0
    for start in range(0, len(lines), text_module.LINES_PER_FRAGMENT):
        chunk = lines[start : start + text_module.LINES_PER_FRAGMENT]
        body = normalise("\n".join(chunk))
        if not body:
            continue
        total += len(body)
        guard_character_budget(total, limits)
        first, last = start + 1, start + len(chunk)
        fragments.append(
            Fragment(
                text=body,
                locator_kind=LocatorKind.LINE_RANGE,
                locator={"line_from": first, "line_to": last},
                locator_label=f"read {first}–{last}",
            )
        )
    if not fragments:
        raise ExtractionFailed("no_text", "Fail on tühi.")
    return ParseResult(
        derivatives=(
            DerivativePayload(
                fragments=tuple(fragments),
                metadata={"encoding": encoding, "line_count": len(lines)},
            ),
        )
    )


def _streamed(content: bytes):
    return parse(TEXT, content)


_BREAKS = ["\n", "\r\n", "\r", "\x0b", "\x0c", "\x1c", "\x1d", "\x1e", "\x85", " ", " "]
_WORDS = ["Koda", "eelnõu", "§ 12", "šaakal", "Žürii", "ÕÄÖÜ", "", "  ", "\t", "a  b", "x"]


def _random_text(rng: random.Random) -> str:
    parts = []
    for _ in range(rng.randrange(0, 900)):
        parts.append(" ".join(rng.choice(_WORDS) for _ in range(rng.randrange(0, 6))))
        parts.append(rng.choice(_BREAKS))
    if rng.random() < 0.5 and parts:
        parts.pop()  # a last line with no break
    return "".join(parts)


def _random_files(count: int) -> list[bytes]:
    rng = random.Random(20261022)  # noqa: S311 - synthetic test text, not a secret
    files = [b"", b"\n", b"\r", b"\r\n", b"a", b"a\r", b"\xef\xbb\xbfBOM\r\nrida"]
    for index in range(count):
        text = _random_text(rng)
        encoding = ("utf-8", "utf-8-sig", "cp1257")[index % 3]
        try:
            files.append(text.encode(encoding))
        except UnicodeEncodeError:
            files.append(text.encode("utf-8"))
    # Undecodable as UTF-8 and as CP1257, and mostly control characters.
    files.append(b"\x98\x81 not text")
    files.append(b"\x01\x02\x03\x04 a")
    return files


@pytest.mark.parametrize("chunk", [1, 2, 3, 7, 64, 1024 * 1024])
def test_streamed_text_reads_exactly_as_the_whole_file_parser(monkeypatch, chunk):
    """Chunk boundaries anywhere — inside "\\r\\n", inside a UTF-8 sequence."""
    monkeypatch.setattr(text_module, "TEXT_CHUNK_BYTES", chunk)
    files = _random_files(60 if chunk > 3 else 12)
    for content in files:
        assert outcome(_streamed, content) == outcome(reference_text_parse, content), content[:80]


def test_text_over_the_character_budget_is_refused_as_before(settings):
    settings.EXTRACTION_MAX_CHARACTERS = 5_000
    rng = random.Random(7)  # noqa: S311 - synthetic test text, not a secret
    for _ in range(20):
        content = _random_text(rng).encode("utf-8") * 20
        assert outcome(_streamed, content) == outcome(reference_text_parse, content)


def test_one_enormous_line_is_refused_before_it_is_whole(settings, monkeypatch):
    settings.EXTRACTION_MAX_CHARACTERS = 10_000
    monkeypatch.setattr(text_module, "TEXT_CHUNK_BYTES", 4096)
    finished: list[int] = []
    original = text_module._PendingLine.finish

    def finish(self):
        line = original(self)
        finished.append(len(line))
        return line

    monkeypatch.setattr(text_module._PendingLine, "finish", finish)
    refused_with("character_limit", _streamed, b"sona " * 1_000_000)
    assert finished == []


def test_a_line_of_spaces_still_reads_as_an_empty_line(monkeypatch):
    monkeypatch.setattr(text_module, "TEXT_CHUNK_BYTES", 4096)
    content = b" " * 300_000 + b"\nteine rida\n"
    assert outcome(_streamed, content) == outcome(reference_text_parse, content)


def test_text_past_the_line_budget_is_refused(settings):
    settings.EXTRACTION_MAX_TEXT_LINES = 1_000
    refused_with("text_line_limit", _streamed, b"rida\n" * 1_001)


def test_text_over_the_size_limit_is_refused_before_it_is_decoded(settings, monkeypatch):
    settings.EXTRACTION_MAX_TEXT_BYTES = 1_000

    def must_not_decode(*args, **kwargs):
        raise AssertionError("decoded")

    monkeypatch.setattr(codecs, "getincrementaldecoder", must_not_decode)
    refused_with("text_size_limit", _streamed, b"a" * 1_001)


def test_text_is_never_decoded_or_split_whole(monkeypatch):
    """A million empty lines: the old parser built a million-entry list."""
    monkeypatch.setattr(text_module, "TEXT_CHUNK_BYTES", 64 * 1024)
    content = b"\n" * 1_000_000
    tracemalloc.start()
    try:
        refused_with("no_text", _streamed, content)
        _current, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert peak < 4 * 1024 * 1024, peak


# ---------------------------------------------------------------------------
# DOCX — csf_eb1ed7815b093aba8872620b
# ---------------------------------------------------------------------------


def reference_docx_parse(content: bytes):
    """The 9a2e834a DocxParser, python-docx's own walk included."""
    import docx

    from app.documents.enums import LocatorKind
    from app.documents.extraction.base import DerivativePayload, Fragment, ParseResult
    from app.documents.extraction.limits import guard_character_budget, guard_office_container

    limits = current_limits()
    handle = io.BytesIO(content)
    guard_office_container(handle, limits)
    try:
        document = docx.Document(handle)
    except Exception as error:
        raise ExtractionFailed("unreadable_docx", "x") from error
    fragments = []
    total = 0
    paragraphs = [normalise(p.text) for p in document.paragraphs]
    first_index = 1
    buffer: list[str] = []
    for index, text in enumerate(paragraphs, start=1):
        if text:
            buffer.append(text)
        if len(buffer) >= 40 or index == len(paragraphs):
            if buffer:
                body = "\n".join(buffer)
                total += len(body)
                guard_character_budget(total, limits)
                fragments.append(
                    Fragment(
                        text=body,
                        locator_kind=LocatorKind.SECTION,
                        locator={"paragraph_from": first_index, "paragraph_to": index},
                        locator_label=f"lõigud {first_index}–{index}",
                    )
                )
                buffer = []
            first_index = index + 1
    for number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [normalise(cell.text) for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                rows.append(line)
        body = "\n".join(rows)
        if not body:
            continue
        total += len(body)
        guard_character_budget(total, limits)
        fragments.append(
            Fragment(
                text=body,
                locator_kind=LocatorKind.SECTION,
                locator={"table": number},
                locator_label=f"tabel {number}",
            )
        )
    for number, section in enumerate(document.sections, start=1):
        for label, part in (("päis", section.header), ("jalus", section.footer)):
            body = normalise("\n".join(p.text for p in part.paragraphs))
            if not body:
                continue
            total += len(body)
            guard_character_budget(total, limits)
            fragments.append(
                Fragment(
                    text=body,
                    locator_kind=LocatorKind.SECTION,
                    locator={"section": number, "part": label},
                    locator_label=f"{label} (osa {number})",
                )
            )
    if not fragments:
        raise ExtractionFailed("no_text", "x")
    return ParseResult(
        derivatives=(
            DerivativePayload(
                fragments=tuple(fragments),
                metadata={
                    "paragraph_count": len(paragraphs),
                    "table_count": len(document.tables),
                    "section_count": len(document.sections),
                    "pagination": "not-available-in-format",
                },
            ),
        )
    )


def _docx(build) -> bytes:
    from docx import Document

    document = Document()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _rich(document) -> None:
    """Paragraph blocks, merged table cells, and sections with linked headers."""
    from docx.enum.section import WD_SECTION

    for number in range(95):
        document.add_paragraph("" if number % 7 == 0 else f"Lõige {number}  koos   tühikutega")
    table = document.add_table(rows=5, cols=4)
    for row in range(5):
        for column in range(4):
            table.cell(row, column).text = f"r{row}c{column}" if (row + column) % 3 else ""
    table.cell(0, 0).merge(table.cell(0, 2))  # a horizontal span
    table.cell(1, 1).merge(table.cell(3, 1))  # a vertical merge
    table.cell(2, 2).merge(table.cell(4, 3))  # both
    document.add_table(rows=2, cols=2)  # an empty one
    document.sections[0].header.paragraphs[0].text = "Viide 1-2/2026"
    for number in range(4):
        section = document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph(f"Osa {number}")
        if number == 1:
            section.header.is_linked_to_previous = False
            section.header.paragraphs[0].text = "Lisa päis"
        if number == 2:
            section.footer.is_linked_to_previous = False
            section.footer.paragraphs[0].text = "Jalus"


def _no_header_first(document) -> None:
    from docx.enum.section import WD_SECTION

    document.add_paragraph("Üks")
    document.add_section(WD_SECTION.NEW_PAGE)
    document.add_paragraph("Kaks")
    section = document.add_section(WD_SECTION.NEW_PAGE)
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = "Ainult kolmandal"


@pytest.mark.parametrize("build", [_rich, _no_header_first], ids=["rich", "first-has-none"])
def test_a_docx_reads_exactly_as_the_python_docx_walk(build):
    content = _docx(build)
    assert outcome(parse, DOCX, content) == outcome(reference_docx_parse, content)


def test_the_synthetic_corpus_docx_reads_exactly_as_before():
    from tests import synthetic_corpus as corpus

    content = corpus.draft_docx()
    assert outcome(parse, DOCX, content) == outcome(reference_docx_parse, content)


def test_docx_tables_over_the_character_budget_are_refused_as_before(settings):
    settings.EXTRACTION_MAX_CHARACTERS = 400
    content = _docx(_rich)
    assert outcome(parse, DOCX, content) == outcome(reference_docx_parse, content)


@pytest.fixture
def no_docx_open(monkeypatch):
    """Call it once the test file is built: python-docx may not open it after."""
    import docx

    def opened(*args, **kwargs):
        raise AssertionError("python-docx opened the package")

    return lambda: monkeypatch.setattr(docx, "Document", opened)


#: python-docx's own template carries styles parts of about 440 KB; a body of
#: 3 000 long paragraphs is a document part well above that, so a part budget
#: between the two is crossed by the document part alone.
BIG_BODY_PART_LIMIT = 500_000


def _big_body(document) -> None:
    for _ in range(3_000):
        document.add_paragraph("x" * 200)


def test_a_docx_part_over_its_budget_is_refused_before_python_docx_parses_it(
    settings, no_docx_open
):
    content = _docx(_big_body)
    settings.EXTRACTION_MAX_DOCX_PART_BYTES = BIG_BODY_PART_LIMIT
    no_docx_open()
    error = refused_with("docx_part_limit", parse, DOCX, content)
    assert "lubatud on 500000" in error.detail


def test_the_docx_xml_total_is_budgeted_across_parts(settings, no_docx_open):
    content = _docx(_rich)
    settings.EXTRACTION_MAX_DOCX_XML_BYTES = BIG_BODY_PART_LIMIT
    no_docx_open()
    refused_with("docx_xml_limit", parse, DOCX, content)


def _rewritten(content: bytes, rewrite) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(content))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            name, data = rewrite(info.filename, source.read(info.filename))
            target.writestr(name, data)
    return out.getvalue()


def test_a_part_renamed_to_look_binary_is_still_charged_as_xml(settings, no_docx_open):
    """The budget follows the declared content type, not the file name."""

    def rename(name: str, data: bytes):
        if name == "word/document.xml":
            return "word/document.png", data
        if name == "word/_rels/document.xml.rels":
            return name, data
        if name == "_rels/.rels":
            return name, data.replace(b"word/document.xml", b"word/document.png")
        if name == "[Content_Types].xml":
            return name, data.replace(b"/word/document.xml", b"/word/document.png")
        return name, data

    content = _rewritten(_docx(_big_body), rename)
    settings.EXTRACTION_MAX_DOCX_PART_BYTES = BIG_BODY_PART_LIMIT
    no_docx_open()
    refused_with("docx_part_limit", parse, DOCX, content)


def test_paragraphs_past_the_budget_are_refused_before_one_is_built(settings, monkeypatch):
    from docx.text import paragraph

    content = _docx(lambda document: [document.add_paragraph("") for _ in range(300)])
    settings.EXTRACTION_MAX_DOCX_PARAGRAPHS = 100
    built = []
    original = paragraph.Paragraph.__init__

    def counting(self, *args, **kwargs):
        built.append(1)
        original(self, *args, **kwargs)

    monkeypatch.setattr(paragraph.Paragraph, "__init__", counting)
    refused_with("docx_paragraph_limit", parse, DOCX, content)
    assert built == []


def _with_document_xml(content: bytes, edit) -> bytes:
    return _rewritten(
        content,
        lambda name, data: (name, edit(data)) if name == "word/document.xml" else (name, data),
    )


def test_a_cell_declaring_an_enormous_span_is_refused_before_the_row_is_expanded(monkeypatch):
    """One XML attribute used to be a million cell objects."""
    from docx.table import _Row

    def expanded(self):
        raise AssertionError("python-docx expanded the row")

    monkeypatch.setattr(_Row, "cells", property(expanded))
    content = _docx(lambda document: document.add_table(rows=1, cols=1))
    content = _with_document_xml(
        content,
        lambda data: data.replace(b"<w:tcPr>", b'<w:tcPr><w:gridSpan w:val="1000000"/>', 1),
    )
    refused_with("table_column_limit", parse, DOCX, content)


@pytest.fixture
def shallow_stack():
    """A recursion limit low enough that a recursive walk is caught quickly."""
    before = sys.getrecursionlimit()
    sys.setrecursionlimit(400)
    try:
        yield 400
    finally:
        sys.setrecursionlimit(before)


def test_a_long_vertical_merge_reads_without_recursion(shallow_stack):
    """The old walk recursed once per merged row and ran out of stack."""
    rows = shallow_stack + 200

    def build(document):
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Ülemine"

    content = _docx(build)
    continue_row = b"<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc></w:tr>"
    content = _with_document_xml(
        content, lambda data: data.replace(b"</w:tbl>", continue_row * rows + b"</w:tbl>", 1)
    )

    assert outcome(reference_docx_parse, content)[0] == "crashed"
    result = parse(DOCX, content)
    [table] = [
        fragment
        for payload in result.derivatives
        for fragment in payload.fragments
        if fragment.locator.get("table") == 1
    ]
    assert table.text == "\n".join(["Ülemine"] * (rows + 1))


def test_many_sections_without_a_header_read_without_recursion(shallow_stack):
    """python-docx resolved each inherited header by recursion, and wrote one in."""
    sections = shallow_stack + 200

    def build(document):
        document.add_paragraph("Algus")

    content = _docx(build)
    break_paragraph = b"<w:p><w:pPr><w:sectPr/></w:pPr><w:r><w:t>osa</w:t></w:r></w:p>"
    content = _with_document_xml(
        content,
        lambda data: data.replace(b"<w:sectPr", break_paragraph * sections + b"<w:sectPr", 1),
    )

    assert outcome(reference_docx_parse, content)[0] == "crashed"
    result = parse(DOCX, content)
    assert result.derivatives[0].metadata["section_count"] == sections + 1


def test_the_docx_walk_is_charged_to_the_table_budget(settings):
    settings.EXTRACTION_MAX_TABLE_CELLS = 10
    content = _docx(lambda document: document.add_table(rows=4, cols=4))
    refused_with("table_cell_limit", parse, DOCX, content)


def test_the_office_module_keeps_one_parser_per_format():
    assert registry.for_mime_type(DOCX) is not None
    assert isinstance(registry.for_mime_type(DOCX), office_module.DocxParser)
