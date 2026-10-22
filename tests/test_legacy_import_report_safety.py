"""The legacy import's report directory: what may be published, and what may run.

Two findings, one directory.

csf_31caadb347feb9939db2cd4e: ``summary.json`` and ``summary.md`` were called
publishable and carried the workbook's file name, its sheet names, duplicate
references, header cells inside the structural findings and the vocabulary
terms it used. They are now written from `PublicSummary` alone — counts,
versions, flags — and the investigative detail moved to the local-only
``operator-summary.*``. The test plants a distinct sentinel in every
source-derived place a workbook has and looks for each of them, recursively,
in both public files.

csf_1cecc4b5be8fb1a7be4fe125: the row-level CSVs wrote workbook cells straight
into files people open in Excel. Every cell, header included, now goes through
`app.core.csv_safety`.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.core.csv_safety import csv_safe
from app.legacy_import import reporting
from app.legacy_import.enums import Anomaly, RowOutcome
from app.legacy_import.inspection import inspect_workbook, write_inspection_reports
from app.legacy_import.reporting import (
    AGGREGATE_REPORTS,
    LOCAL_ONLY_REPORTS,
    OPERATOR_REPORTS,
    PublicSummary,
    PublicSummaryError,
)
from tests.synthetic_register import Row, Sheet, write_workbook

#: One distinct, unguessable word per place a workbook can put text. Each is
#: unlike every Estonian word and every code this application writes, so a hit
#: is a leak and never a coincidence.
SENTINELS = {
    "file_name": "Qzfailinimiqz",
    "sheet_name": "Qzlehenimiqz",
    "mismatched_header": "Qzpealkirjalahterqz",
    "uncontracted_header": "Qzlisaveeruqz",
    "uncontracted_cell": "Qzlisalahterqz",
    "title": "Qzteemapealkiriqz",
    "counterparty": "Qzasutusqz",
    "owner": "Qzvastutajaqz",
    "status": "Qzhetkeseisqz",
    "next_action": "Qzjargmisekstqz",
    "hyperlink": "qzlinkqz.invalid",
    "vocabulary_label": "Qzsonavaraqz",
    "vocabulary_explanation": "Qzselgitusqz",
}

DUPLICATE_REFERENCE = "2026_4271"

#: A row whose every text cell is a formula a spreadsheet would run.
FORMULA_ROW = Row(
    reference="2026_3",
    title='=HYPERLINK("http://paha.invalid","Vaata")',
    counterparty="@SUM(1+1)*cmd",
    owner="+SUM(A1:A9)",
    status="\t=1+1",
    next_action="-cmd|' /c calc'!A0",
)


def _sentinel_workbook(directory: Path) -> Path:
    path = write_workbook(
        directory / f"{SENTINELS['file_name']}.xlsx",
        [
            Sheet(
                2026,
                [
                    Row(
                        reference=DUPLICATE_REFERENCE,
                        title=SENTINELS["title"],
                        counterparty=SENTINELS["counterparty"],
                        owner=SENTINELS["owner"],
                        status=SENTINELS["status"],
                        next_action=SENTINELS["next_action"],
                        hyperlink=f"onenote:https://{SENTINELS['hyperlink']}/leht",
                    ),
                    Row(reference=DUPLICATE_REFERENCE, title="Sünteetiline teine", owner="Kadri"),
                    FORMULA_ROW,
                ],
            ),
            Sheet(2025, [Row(reference="2025_1", title="Sünteetiline 2025", owner="Kadri")]),
        ],
    )
    workbook = load_workbook(path)
    # A sheet that is neither a year nor the vocabulary.
    workbook.create_sheet(SENTINELS["sheet_name"]).append(["midagi"])
    # A contracted header that says something else: a structural finding that
    # quotes the cell.
    workbook["2025"].cell(row=1, column=2).value = SENTINELS["mismatched_header"]
    # A labelled column the contract does not describe, holding data.
    sheet_2026 = workbook["2026"]
    sheet_2026.cell(row=1, column=20).value = SENTINELS["uncontracted_header"]
    sheet_2026.cell(row=2, column=20).value = SENTINELS["uncontracted_cell"]
    # openpyxl stores a string beginning with `=` as a formula, which a
    # data-only read returns as nothing. A register cell holds such text as
    # text, so the formula row's cells are written back as strings.
    for column, typed in (
        (2, FORMULA_ROW.title),
        (7, FORMULA_ROW.counterparty),
        (8, FORMULA_ROW.owner),
        (11, FORMULA_ROW.status),
        (12, FORMULA_ROW.next_action),
    ):
        cell = sheet_2026.cell(row=4, column=column)
        cell.value = typed
        cell.data_type = "s"
    # The workbook's own vocabulary sheet, with a term the seed does not know.
    workbook["Hetkeseisu info"].append(
        [SENTINELS["vocabulary_label"], SENTINELS["vocabulary_explanation"]]
    )
    workbook.save(path)
    workbook.close()
    return path


def _strings(value: Any) -> list[str]:
    """Every key and every value in a JSON document, as text, recursively."""
    if isinstance(value, dict):
        return [text for key, item in value.items() for text in (str(key), *_strings(item))]
    if isinstance(value, list):
        return [text for item in value for text in _strings(item)]
    return [str(value)]


def _assert_public_carries_none_of(directory: Path, forbidden: list[str]) -> None:
    payload = json.loads((directory / "summary.json").read_text(encoding="utf-8"))
    PublicSummary.from_payload(payload)  # the file is exactly the schema
    strings = [text.casefold() for text in _strings(payload)]
    texts = {
        name: (directory / name).read_text(encoding="utf-8").casefold()
        for name in AGGREGATE_REPORTS
    }
    for word in forbidden:
        needle = word.casefold()
        assert not any(needle in text for text in strings), f"{word!r} in summary.json"
        for name, text in texts.items():
            assert needle not in text, f"{word!r} in {name}"


@pytest.fixture
def inspected(tmp_path: Path):
    path = _sentinel_workbook(tmp_path / "source")
    directory = tmp_path / "reports"
    written = write_inspection_reports(inspect_workbook(path), directory)
    return directory, written


# ---------------------------------------------------------------------------
# csf_31caadb347feb9939db2cd4e — the public summaries
# ---------------------------------------------------------------------------


def test_no_sentinel_reaches_either_public_summary(inspected):
    directory, _written = inspected
    operator = json.loads((directory / "operator-summary.json").read_text(encoding="utf-8"))
    duplicates = operator["duplicate_references"]
    assert duplicates, "the fixture planted no duplicate reference"

    _assert_public_carries_none_of(
        directory, [*SENTINELS.values(), *duplicates, DUPLICATE_REFERENCE, "paha.invalid"]
    )


def test_the_public_summary_is_only_counts_versions_and_flags(inspected):
    directory, _written = inspected
    payload = json.loads((directory / "summary.json").read_text(encoding="utf-8"))

    assert set(payload) == {
        "schema",
        "schema_version",
        "report",
        "parser_version",
        "importer_version",
        "contract_versions",
        "counts",
        "categories",
        "flags",
    }
    assert set(payload["counts"]) <= reporting.PUBLIC_COUNTS
    assert all(type(value) is int for value in payload["counts"].values())
    assert set(payload["flags"]) <= reporting.PUBLIC_FLAGS
    assert all(type(value) is bool for value in payload["flags"].values())
    assert set(payload["categories"]["anomalies"]) <= set(Anomaly.values)
    assert set(payload["categories"]["outcomes"]) <= set(RowOutcome.values)
    # The structure the sentinels created is still *counted*, just not quoted.
    assert payload["counts"]["structure_findings"] >= 2
    assert payload["counts"]["duplicate_references"] == 1
    assert payload["counts"]["vocabulary_labels_missing_from_seed"] >= 1
    assert payload["flags"]["structure_ok"] is False


def test_the_operator_summary_keeps_the_investigative_detail(inspected):
    """Local-only files still say everything a reviewer needs to act."""
    directory, written = inspected
    assert {path.name for path in written.local_only} == set(LOCAL_ONLY_REPORTS)
    assert {path.name for path in written.operator} == set(OPERATOR_REPORTS)
    assert {path.name for path in written.aggregate} == set(AGGREGATE_REPORTS)

    operator = (directory / "operator-summary.json").read_text(encoding="utf-8")
    operator_md = (directory / "operator-summary.md").read_text(encoding="utf-8")
    for key in (
        "file_name",
        "sheet_name",
        "mismatched_header",
        "uncontracted_header",
        "status",
        "vocabulary_label",
    ):
        assert SENTINELS[key] in operator, key
    assert SENTINELS["file_name"] in operator_md
    assert SENTINELS["mismatched_header"] in operator_md
    assert "AINULT KOHALIK" in operator_md

    rows = (directory / "rows.csv").read_text(encoding="utf-8-sig")
    for key in ("title", "counterparty", "owner", "next_action", "hyperlink"):
        assert SENTINELS[key] in rows, key


def _payload(**overrides: Any) -> dict[str, Any]:
    payload = PublicSummary(
        report="inspection",
        parser_version="1.0",
        contract_versions="1.1/2011-2026",
        counts={"matter_rows": 3},
        categories={"anomalies": {Anomaly.DUPLICATE_REFERENCE.value: 2}},
        flags={"structure_ok": True},
    ).to_payload()
    payload.update(overrides)
    return payload


def test_the_schema_round_trips():
    payload = _payload()
    assert PublicSummary.from_payload(payload).to_payload() == payload


@pytest.mark.parametrize(
    ("overrides", "refusal"),
    [
        ({"source": {"file_name": "x.xlsx"}}, "ei luba välju"),
        ({"duplicate_references": ["2026_1"]}, "ei luba välju"),
        ({"schema": "something-else"}, "skeem ei ole tuttav"),
        ({"schema_version": 2}, "skeem ei ole tuttav"),
        ({"report": "Qzaruanneqz"}, "liik ei ole lubatud"),
        ({"parser_version": "Kordumatu pealkiri"}, "ei ole versioon"),
        ({"counts": {"sheet_names": 1}}, "ei luba loendeid"),
        ({"counts": {"matter_rows": "3"}}, "ei ole loend"),
        ({"counts": {"matter_rows": True}}, "ei ole loend"),
        ({"counts": {"matter_rows": -1}}, "ei ole loend"),
        ({"flags": {"has_titles": True}}, "ei luba märke"),
        ({"flags": {"structure_ok": 1}}, "ei ole jah/ei"),
        ({"categories": {"vocabulary": {}}}, "ei luba jaotusi"),
        ({"categories": {"anomalies": {"Qzvotiqz": 1}}}, "ei ole rakenduse kood"),
    ],
)
def test_the_schema_refuses_what_it_does_not_name(overrides, refusal):
    with pytest.raises(PublicSummaryError, match=refusal) as refused:
        PublicSummary.from_payload(_payload(**overrides))
    # A refused source string is not repeated back in the refusal.
    assert "Qzvotiqz" not in str(refused.value)


def test_the_public_files_cannot_be_written_from_a_dictionary(tmp_path: Path):
    with pytest.raises(PublicSummaryError, match="peab olema PublicSummary"):
        reporting.write_reports(
            tmp_path,
            public={"source": {"file_name": "x.xlsx"}},  # type: ignore[arg-type]
            operator_summary={},
            operator_markdown="",
            rows=[],
            row_header=[],
            anomalies=[],
            anomaly_header=[],
            mapping_gaps=[],
            mapping_gap_header=[],
            candidates=[],
            candidate_header=[],
        )
    assert not (tmp_path / "summary.json").exists()


@pytest.mark.django_db
def test_the_plan_reports_publish_no_sentinel_either(tmp_path: Path):
    """The database-aware run writes the same three classes through the same schema."""
    from django.core.management import call_command

    from tests import factories

    factories.UserFactory(display_name="Kadri", upn="kadri@example.invalid")
    path = _sentinel_workbook(tmp_path / "source")
    directory = tmp_path / "plan"
    call_command("import_legacy_register", str(path), "--dry-run", "--report-dir", str(directory))

    payload = json.loads((directory / "summary.json").read_text(encoding="utf-8"))
    assert payload["report"] == "dry-run"
    assert "accounting_is_complete" in payload["flags"]
    operator = (directory / "operator-summary.json").read_text(encoding="utf-8")
    assert SENTINELS["file_name"] in operator
    assert SENTINELS["mismatched_header"] in operator

    _assert_public_carries_none_of(
        directory, [*SENTINELS.values(), DUPLICATE_REFERENCE, "paha.invalid"]
    )
    _assert_every_csv_cell_is_inert(directory)


# ---------------------------------------------------------------------------
# csf_1cecc4b5be8fb1a7be4fe125 — the CSV reports
# ---------------------------------------------------------------------------


def _read_csv(path: Path) -> list[list[str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle, delimiter=";"))


def _assert_every_csv_cell_is_inert(directory: Path) -> None:
    """No cell is left that `csv_safe` would still change."""
    for name in reporting.ROW_LEVEL_REPORTS:
        for row in _read_csv(directory / name):
            for cell in row:
                assert csv_safe(cell) == cell, (name, cell)


def test_every_row_level_cell_reaches_excel_as_text(inspected):
    directory, _written = inspected
    _assert_every_csv_cell_is_inert(directory)

    rows = _read_csv(directory / "rows.csv")
    formula = next(row for row in rows if FORMULA_ROW.title in row[4])
    for column, typed in (
        (4, FORMULA_ROW.title),
        (5, FORMULA_ROW.counterparty),
        (7, FORMULA_ROW.owner),
        (13, FORMULA_ROW.status),
        (14, FORMULA_ROW.next_action),
    ):
        assert typed is not None
        cell = formula[column]
        assert cell.startswith("'"), (column, cell)
        assert cell[1:].strip() == typed.strip(), (column, cell)


def test_the_csv_writer_neutralises_headers_and_keeps_numbers(tmp_path: Path):
    path = tmp_path / "x.csv"
    reporting._write_csv(
        path,
        ["=leht", "rida", "\r@viide"],
        [
            [-5, Decimal("-9.5"), dt.date(2026, 9, 12)],
            ["-5", "-3,2", "+SUM(A1)"],
            ["\n=1+1", " =1+1", "tavaline"],
        ],
    )
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf"), "utf-8-sig, so Excel reads the encoding"
    header, numbers, numeric_text, formulas = _read_csv(path)
    assert header == ["'=leht", "rida", "'\r@viide"]
    assert numbers == ["-5", "-9.5", "2026-09-12"]
    assert numeric_text == ["-5", "-3,2", "'+SUM(A1)"]
    assert formulas == ["'\n=1+1", "' =1+1", "tavaline"]


def test_the_owner_backfill_report_is_inert(tmp_path: Path):
    from types import SimpleNamespace

    from app.legacy_import.management.commands.backfill_legacy_owners import Command

    path = tmp_path / "vastutajad.csv"
    plan = SimpleNamespace(unresolved_values={"=HYPERLINK(1)": 2, "Kadri": 1})
    command = Command()
    command._write_unresolved(plan, path)
    rows = _read_csv(path)
    assert rows[0] == ["vastutaja allikas", "esinemisi"]
    assert ["'=HYPERLINK(1)", "2"] in rows
    assert ["Kadri", "1"] in rows
